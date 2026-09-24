"""Generate sample bank statements (Excel, CSV, PDF) for tests and manual trials."""
import os

import openpyxl

ROWS = [
    ("01/04/2026", "NEFT-AWS INDIA-AMAZON WEB SERVICES", "N123", 8500.00, None, 91500.00),
    ("03/04/2026", "UPI/ACME TRADERS/acme@okhdfcbank/Payment", "U456", None, 12000.00, 103500.00),
    ("05/04/2026", "CHQ PAID-OFFICE RENT APRIL", "000781", 25000.00, None, 78500.00),
    ("07/04/2026", "ATM WDL MG ROAD BANGALORE", "", 2000.00, None, 76500.00),
]
HEADER = ["Date", "Narration", "Chq./Ref.No.", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]


def make_excel(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["HDFC BANK LTD"])
    ws.append(["Statement of account", "", "Account No: 50100012345678"])
    ws.append([])
    ws.append(HEADER)
    for r in ROWS:
        ws.append(list(r))
    ws.append([])
    ws.append(["", "STATEMENT SUMMARY", "", "35500.00", "12000.00", ""])
    wb.save(path)
    return path


def make_excel_with_ledger(path):
    """User-prepared sheet with a Ledger column (one wrong case, one unknown, one blank)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date", "Particulars", "Withdrawal", "Deposit", "Balance", "Ledger Name"])
    ws.append(["01/04/2026", "NEFT-AWS INDIA", 8500.00, None, 91500.00, "Amazon Web Services"])
    ws.append(["03/04/2026", "UPI/ACME TRADERS/Payment", None, 12000.00, 103500.00, "acme traders"])
    ws.append(["05/04/2026", "CHQ PAID-OFFICE RENT APRIL", 25000.00, None, 78500.00, "Rent Account"])
    ws.append(["07/04/2026", "ATM WDL MG ROAD", 2000.00, None, 76500.00, ""])
    wb.save(path)
    return path


def make_csv_single_amount(path):
    with open(path, "w") as f:
        f.write("Txn Date,Description,Amount,Dr/Cr,Balance\n")
        f.write("2026-04-01,SALARY CREDIT,50000.00,CR,150000.00\n")
        f.write("2026-04-02,ELECTRICITY BILL,1,234.50,DR,148765.50\n".replace("1,234.50", '"1,234.50"'))
    return path


def make_pdf_table(path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    data = [HEADER] + [[c if c is not None else "" for c in r] for r in ROWS]
    data = [[("%.2f" % c) if isinstance(c, float) else c for c in row] for row in data]
    table = Table(data)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc = SimpleDocTemplate(path, pagesize=landscape(A4))
    doc.build([Paragraph("HDFC BANK - Statement of account", getSampleStyleSheet()["Title"]), table])
    return path


def make_pdf_text(path):
    """A PDF with plain text lines (no ruling), like many bank e-statements."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(path, pagesize=A4)
    y = 800
    for line in ["STATE BANK OF INDIA", "Date Description Debit Credit Balance",
                 "Opening balance 100000.00",
                 "01-04-2026 NEFT AWS INDIA 8,500.00 91,500.00",
                 "03-04-2026 UPI ACME TRADERS 12,000.00 1,03,500.00",
                 "05-04-2026 CHQ OFFICE RENT 25,000.00 78,500.00"]:
        c.drawString(40, y, line)
        y -= 18
    c.save()
    return path


def make_encrypted_pdf(path, password="secret"):
    """Password-protected PDF statement with a ruled table (like most bank e-statements)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    data = [HEADER] + [[("%.2f" % c) if isinstance(c, float) else (c or "") for c in r] for r in ROWS]
    table = Table(data)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(path, pagesize=landscape(A4), encrypt=password).build([table])
    return path


def make_encrypted_excel(path, password="secret"):
    """Excel statement encrypted with a password (File > Protect Workbook > Encrypt with Password)."""
    from msoffcrypto.format.ooxml import OOXMLFile

    plain = path + ".plain.xlsx"
    make_excel(plain)
    with open(plain, "rb") as f, open(path, "wb") as out:
        OOXMLFile(f).encrypt(password, out)
    os.remove(plain)
    return path


def make_disguised_xls(path, kind):
    """Files banks send with an .xls name that are really something else:
    'xlsx' (a modern workbook), 'html' (an HTML table), 'xml2003' (Excel 2003 XML), 'tsv' (text)."""
    header = ["Txn Date", "Description", "Ref No", "Txn Amount", "Balance"]
    rows = [["07/04/2026", "SMS CHARGES", "0", "\u20b9\u00a02.12 Dr", "\u20b9\u00a01,191.12"],
            ["01/07/2026", "INTEREST ON SB A/C", " ", "\u20b9\u00a07.00 Cr", "\u20b9\u00a01,198.12"]]
    letterhead = [["CASA Statement"], ["Customer Name", "A KUMAR"], []]
    summary = [[], ["Opening Balance", "Debit Count", "Total Debits", "Credit Count", "Total Credits"],
               ["\u20b9\u00a01,193.24", "1", "\u20b9\u00a02.12", "1", "\u20b9\u00a07.00"]]
    all_rows = letterhead + [header] + rows + summary
    if kind == "xlsx":
        wb = openpyxl.Workbook()
        for r in all_rows:
            wb.active.append(r)
        wb.save(path)  # openpyxl writes xlsx content whatever the extension
    elif kind == "html":
        cells = lambda r, tag: "".join("<%s>%s</%s>" % (tag, c, tag) for c in r)
        body = "".join("<tr>%s</tr>" % cells(r, "th" if r is header else "td") for r in all_rows)
        with open(path, "w", encoding="utf-8") as f:
            f.write("<html><head><style>td{mso-number-format:'\\@'}</style></head><body>"
                    "<table border=1>%s</table></body></html>" % body)
    elif kind == "xml2003":
        def row_xml(r):
            return "<Row>%s</Row>" % "".join('<Cell><Data ss:Type="String">%s</Data></Cell>' % c for c in r)
        with open(path, "w", encoding="utf-8") as f:
            f.write('<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
                    'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="S"><Table>'
                    + "".join(row_xml(r) for r in all_rows) + "</Table></Worksheet></Workbook>")
    elif kind == "tsv":
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join("\t".join(r) for r in all_rows))
    return path


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
    os.makedirs(out, exist_ok=True)
    make_excel(os.path.join(out, "sample_statement.xlsx"))
    make_csv_single_amount(os.path.join(out, "sample_statement.csv"))
    make_excel_with_ledger(os.path.join(out, "sample_statement_with_ledger.xlsx"))
    make_pdf_table(os.path.join(out, "sample_statement.pdf"))
    make_pdf_text(os.path.join(out, "sample_statement_text.pdf"))
    make_encrypted_pdf(os.path.join(out, "sample_statement_protected.pdf"))      # password: secret
    make_encrypted_excel(os.path.join(out, "sample_statement_protected.xlsx"))   # password: secret
    print("Samples written to", out)
